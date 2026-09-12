#!/usr/bin/env python3
"""Fill the case's submission template (candidate fields + full technical
answer) without disturbing the original template content above it.

Usage:
    python scripts/fill_docx_template.py <template.docx> <output.docx>

The template's own text (the case statement, evaluation criteria, etc.) is
left untouched -- this script only (a) fills the "Nome Completo / E-mail /
CPF" table and (b) appends the candidate's answer as new sections after the
template's final paragraph.
"""

from __future__ import annotations

import os
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

CANDIDATE_NAME = "Enzo Nicolás Spotorno Bieger"
CANDIDATE_EMAIL = "enzonsb@gmail.com"
GITHUB_URL = "https://github.com/enzoniko/predictive-maintenance-senai"

# Never hardcoded (this script is committed to a public repository): read
# from an environment variable at generation time, or leave blank and fill
# it in Word before submitting.
CANDIDATE_CPF = os.environ.get("CANDIDATE_CPF", "")


def fill_candidate_fields(doc: Document) -> None:
    for table in doc.tables:
        header_row = [c.text.strip() for c in table.rows[0].cells]
        if "Nome Completo" in header_row:
            table.rows[0].cells[1].text = CANDIDATE_NAME
            second_row = [c.text.strip() for c in table.rows[1].cells]
            if "E-mail" in second_row:
                table.rows[1].cells[1].text = CANDIDATE_EMAIL
            if "CPF" in second_row:
                table.rows[1].cells[3].text = CANDIDATE_CPF
            return
    print("WARNING: candidate-fields table not found -- fill Nome/E-mail/CPF manually.",
          file=sys.stderr)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_paragraph(doc: Document, text: str, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_lead:
        run = p.add_run(bold_lead)
        run.bold = True
        p.add_run(text)
    else:
        p.add_run(text)


def add_bullets(doc: Document, items: list[str]) -> None:
    # The template doesn't define a "List Bullet" style (it's a stripped-down
    # export), so bullets are drawn as plain characters rather than relying
    # on a style that may not exist in every Word template.
    for item in items:
        doc.add_paragraph(f"•  {item}")


def append_answer(doc: Document) -> None:
    doc.add_page_break()

    title = doc.add_heading("Resolução do Estudo de Caso", level=0)
    for run in title.runs:
        run.font.size = Pt(20)

    add_paragraph(
        doc,
        f"Este documento resume a interpretação do problema, a arquitetura proposta, o "
        f"cronograma de execução e os resultados da prévia técnica. O repositório completo "
        f"(código, testes, notebooks executados, documentação detalhada e apresentações) "
        f"está em {GITHUB_URL}.",
    )

    add_heading(doc, "1. Interpretação do problema", level=1)
    add_paragraph(
        doc,
        "O case exige dois papéis simultâneos: mentor/arquiteto de um projeto de manutenção "
        "preditiva (arquitetura + cronograma, caso contratado) e autor de uma prévia técnica "
        "que convença o diretor a contratar o projeto. Antes de treinar qualquer modelo, os "
        "dados fornecidos foram auditados estatisticamente (não apenas inspecionados): dos "
        "5 arquivos de sensores, 3 (Dados_1-3) carregam sinal real e discriminam as 5 classes "
        "com significância estatística; Dados_4 está travado num valor quase constante "
        "(sensor morto/mal cabeado) e Dados_5 é estatisticamente indistinguível de ruído "
        "branco em relação às classes -- ambos excluídos com teste de Kruskal-Wallis + "
        "permutação de informação mútua, não por inspeção visual. Uma coluna extra presente "
        "em Dados_1-3 (201ª, quase inteiramente vazia) foi identificada como artefato de "
        "exportação e descartada com validação. As 5 classes vêm artificialmente balanceadas "
        "(10.000 janelas cada), padrão que não ocorre numa planta real -- a avaliação final "
        "reamostra o teste sob um prior de classes realista. Detalhes completos: "
        "docs/01_interpretacao_problema.md e docs/07_perguntas_ao_cliente.md (perguntas em "
        "aberto para o cliente).",
    )

    add_heading(doc, "2. Arquitetura proposta", level=1)
    add_paragraph(
        doc,
        "Pipeline: auditoria de dados -> limpeza (filtro de Hampel para saturação, "
        "interpolação de NaN, flag de janela silenciosa) -> engenharia de features (tempo, "
        "frequência, wavelet) -> seleção de modelo por validação cruzada (árvore de decisão, "
        "Random Forest, HistGradientBoosting, XGBoost) -> calibração de predição conformal -> "
        "avaliação completa (calibração, desbalanceamento simulado, robustez, "
        "explicabilidade) -> ModelBundle versionado -> API FastAPI (/predict, /predict_batch, "
        "/explain, /audit, /drift, /health) -> banco de dados (SQLite local / Postgres via "
        "Docker) -> dashboard Streamlit. O banco de dados desta camada de IA armazena apenas "
        "o que a IA produz (predições, auditorias, relatórios de drift) -- não duplica o "
        "histórico bruto de sensores, que permanece no banco já modelado pelo time de "
        "software do cliente. Diagramas completos (contexto, containers, sequência de "
        "inferência): docs/03_arquitetura.md.",
    )
    add_paragraph(
        doc,
        "Desenvolvido em Windows ARM64 e validado continuamente num servidor Linux x86-64 -- "
        "os dois alvos de implantação. Toda dependência que falha numa plataforma "
        "(mlflow, streamlit, shap, ssqueezepy, psycopg2, pyarrow) tem fallback documentado e "
        "testado: ver docs/03_arquitetura.md, seção 3.6.",
    )

    add_heading(doc, "3. Cronograma de execução (caso contratado) -- 9 meses", level=1)
    add_bullets(
        doc,
        [
            "Fase 0 (mês 1) -- Kick-off e requisitos: semântica das classes, unidade física "
            "dos sensores, acesso ao banco de dados do cliente.",
            "Fase 1 (mês 2) -- Entendimento: revisão de literatura, auditoria de dados no "
            "banco real do cliente.",
            "Fase 2 (meses 3-4) -- Baseline interpretável: features físicas + modelo "
            "auditável, validado com o time de manutenção.",
            "Fase 3 (meses 4-6, paralelo) -- Pesquisa avançada: physics-informed, "
            "open-set/não-supervisionado, predição conformal em produção.",
            "Fase 4 (meses 6-7) -- Integração: API + banco + dashboard, piloto numa máquina.",
            "Fase 5 (mês 8) -- Validação de campo e transferência de conhecimento.",
            "Fase 6 (mês 9+) -- Acompanhamento contínuo: drift, retrain, evolução.",
        ],
    )
    add_paragraph(
        doc,
        "WBS, esforço estimado por fase e Gantt completo: docs/04_cronograma_9_meses.md.",
    )

    add_heading(doc, "4. Riscos e track de pesquisa", level=1)
    add_paragraph(
        doc,
        "Dez riscos priorizados (dados insuficientes, sensores em falha, desbalanceamento "
        "real, falhas nunca vistas, drift de conceito, resistência de adoção, overfitting ao "
        "conjunto curado, escala, portabilidade de dependências, disponibilidade do cliente) "
        "-- a maioria já com o mecanismo de mitigação prototipado no código, não apenas "
        "planejado. Ver docs/05_riscos_e_mitigacoes.md.",
    )
    add_paragraph(
        doc,
        "O track de pesquisa (docs/06_track_pesquisa.md) formaliza 8 hipóteses -- modelos "
        "informados por física, aprendizado não/semi/auto-supervisionado para o regime real "
        "de poucos dados de falha, reconhecimento open-set e predição conformal recalibrada "
        "em produção -- fundamentadas em publicações do autor sobre diagnóstico de máquinas "
        "rotativas (IEEE IECON 2025, IEEE Access 2026), com datasets públicos (CWRU, MFPT, "
        "Paderborn, MaFaulDa) propostos para viabilizar a pesquisa em paralelo caso os dados "
        "do cliente demorem a ficar prontos.",
    )

    add_heading(doc, "5. Resultados da prévia técnica", level=1)
    add_paragraph(
        doc,
        "Modelo final (HistGradientBoosting, vencedor da validação cruzada entre 4 "
        "candidatos, nas duas plataformas testadas): execução completa no servidor Linux "
        "x86-64 com a stack de pesquisa inteira (MLflow, SHAP, ssqueezepy) atinge "
        "F1-macro = 0,962 em teste nunca visto, erro de calibração esperado (ECE) = 0,006, "
        "cobertura de predição conformal = 0,894 contra um alvo configurado de 0,90 "
        "(tamanho médio do conjunto: 0,98 -- quase sempre uma única classe confiante). "
        "Dois controles de sanidade confirmam ausência de vazamento: rótulos embaralhados "
        "ficam em F1 = 0,200 (exatamente o acaso teórico); um modelo treinado só no sensor "
        "de ruído excluído fica em F1 = 0,067, abaixo até do acaso. Testes de robustez "
        "(remoção de sensor, remoção de grupo de features) e explicabilidade (SHAP real) "
        "completam a avaliação. Reprodutível via `python -m pdm.cli train` e detalhado "
        "nos notebooks 01-03 e em docs/01_interpretacao_problema.md e "
        "docs/03_arquitetura.md (seção 3.7, validação cruzada de plataforma).",
    )

    add_heading(doc, "6. Links", level=1)
    add_bullets(
        doc,
        [
            f"Repositório no GitHub: {GITHUB_URL}",
            "Vídeos (roteiros em docs/videos/; links dos vídeos gravados: preencher antes do envio)",
            "Apresentação executiva e técnica: anexos separados deste envio",
        ],
    )


def main() -> int:
    if len(sys.argv) != 3:
        print(__doc__)
        return 1
    template_path, out_path = Path(sys.argv[1]), Path(sys.argv[2])

    doc = Document(str(template_path))
    fill_candidate_fields(doc)
    append_answer(doc)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"Wrote {out_path}")
    if not CANDIDATE_CPF:
        print("REMINDER: CPF field left blank -- fill it in Word before submitting.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
